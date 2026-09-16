"""load_sharepoint_kb_folder: one bad file (wrong format, malformed headings)
must never block the rest of the folder from indexing."""

import io

from docx import Document

from app.services.sharepoint_kb_loader import load_sharepoint_kb_folder


def _docx_bytes(headings_and_bodies: list[tuple[int, str, list[str]]]) -> bytes:
    doc = Document()
    for level, heading, body in headings_and_bodies:
        doc.add_heading(heading, level=level)
        for line in body:
            doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


VALID_DOCX = _docx_bytes([(1, "How can I reschedule my exam?", ["Send us a request."])])
EMPTY_DOCX = _docx_bytes([])  # no headings at all


class _FakeClient:
    def __init__(self, items, files, drives=None):
        self._items = items          # what list_drive_items returns
        self._files = files          # item_id -> bytes
        self._drives = drives or {}  # library name -> drive id

    def get_site(self, hostname, site_path):
        return {"id": "site1"}

    def get_drive_id(self, site_id, library_name):
        return self._drives[library_name]

    def list_drive_items(self, site_id, path="", drive_id=None):
        return {"value": self._items}

    def download_file(self, site_id, item_id, drive_id=None):
        return self._files[item_id]


def test_happy_path_indexes_every_docx():
    client = _FakeClient(
        items=[
            {"id": "1", "name": "Scheduling.docx"},
            {"id": "2", "name": "Technical.docx"},
        ],
        files={"1": VALID_DOCX, "2": VALID_DOCX},
    )
    documents, warnings = load_sharepoint_kb_folder(
        client, "host", "/sites/x", "General FAQ"
    )
    assert warnings == []
    assert [d[0] for d in documents] == ["Scheduling.docx", "Technical.docx"]


def test_folders_in_the_listing_are_skipped_not_indexed():
    client = _FakeClient(
        items=[
            {"id": "1", "name": "Scheduling.docx"},
            {"id": "2", "name": "Archive", "folder": {"childCount": 3}},
        ],
        files={"1": VALID_DOCX},
    )
    documents, warnings = load_sharepoint_kb_folder(client, "host", "/x", "General")
    assert [d[0] for d in documents] == ["Scheduling.docx"]
    assert warnings == []


def test_non_docx_file_is_skipped_with_a_warning_others_still_index():
    client = _FakeClient(
        items=[
            {"id": "1", "name": "Scheduling.docx"},
            {"id": "2", "name": "old_notes.pdf"},
        ],
        files={"1": VALID_DOCX, "2": b"%PDF-1.4 not real content"},
    )
    documents, warnings = load_sharepoint_kb_folder(client, "host", "/x", "General")
    assert [d[0] for d in documents] == ["Scheduling.docx"]
    assert len(warnings) == 1
    assert "old_notes.pdf" in warnings[0]
    assert ".docx" in warnings[0]


def test_a_malformed_docx_is_skipped_others_still_index():
    client = _FakeClient(
        items=[
            {"id": "1", "name": "Scheduling.docx"},
            {"id": "2", "name": "Corrupt.docx"},
        ],
        files={"1": VALID_DOCX, "2": b"this is not a real docx file"},
    )
    documents, warnings = load_sharepoint_kb_folder(client, "host", "/x", "General")
    assert [d[0] for d in documents] == ["Scheduling.docx"]
    assert any("Corrupt.docx" in w for w in warnings)


def test_a_docx_with_no_usable_headings_is_skipped():
    client = _FakeClient(
        items=[{"id": "1", "name": "Empty.docx"}],
        files={"1": EMPTY_DOCX},
    )
    documents, warnings = load_sharepoint_kb_folder(client, "host", "/x", "General")
    assert documents == []
    # Both docx_to_markdown's own "no headings" warning AND the loader's
    # "nothing usable" warning should surface, prefixed with the filename.
    assert any("Empty.docx" in w for w in warnings)


def test_library_name_resolves_the_named_drive_and_reads_from_it():
    """A KB kept in its own (non-default) library must actually be read from
    that library, not silently fall back to the default one."""
    client = _FakeClient(
        items=[{"id": "1", "name": "Scheduling.docx"}],
        files={"1": VALID_DOCX},
        drives={"Candidate experience KB": "drive-kb"},
    )
    documents, warnings = load_sharepoint_kb_folder(
        client, "host", "/sites/everybody", "General",
        library_name="Candidate experience KB",
    )
    assert warnings == []
    assert [d[0] for d in documents] == ["Scheduling.docx"]


def test_no_library_name_never_calls_get_drive_id():
    """Every existing caller omits library_name, so get_drive_id must not be
    invoked at all in that path — a client without a real per-site library
    must keep working exactly as it did before this feature existed."""
    client = _FakeClient(
        items=[{"id": "1", "name": "Scheduling.docx"}],
        files={"1": VALID_DOCX},
        drives={},  # get_drive_id would KeyError if called — proves it wasn't
    )
    documents, _ = load_sharepoint_kb_folder(client, "host", "/x", "General")
    assert [d[0] for d in documents] == ["Scheduling.docx"]


def test_per_heading_warnings_are_prefixed_with_the_filename():
    """A heading with no answer inside an otherwise-valid file still
    produces a warning naming which FILE it came from — important once
    there are many files in a folder."""
    doc = Document()
    doc.add_heading("Scheduling", level=1)  # bare section title, no answer
    doc.add_heading("How can I reschedule my exam?", level=1)
    doc.add_paragraph("Send us a request.")
    buf = io.BytesIO()
    doc.save(buf)

    client = _FakeClient(
        items=[{"id": "1", "name": "Scheduling.docx"}],
        files={"1": buf.getvalue()},
    )
    documents, warnings = load_sharepoint_kb_folder(client, "host", "/x", "General")
    assert [d[0] for d in documents] == ["Scheduling.docx"]
    assert len(warnings) == 1
    assert warnings[0].startswith("Scheduling.docx: ")
    assert "dropped" in warnings[0]
