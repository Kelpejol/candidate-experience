"""Tests for the Word-document KB authoring standard: .docx only, Heading 1/2
only, one heading = one question, no bare section titles.

Builds real .docx files in-memory with python-docx's own Document API (not
hand-crafted bytes), so these tests exercise the actual file format, not a
guess at its shape.
"""

import io

import pytest
from docx import Document

from app.services.docx_kb_loader import docx_to_markdown
from app.services.voice_kb_ingest import chunk_markdown


def _build_docx(*, headings_and_bodies: list[tuple[int, str, list[str] | None]] = None,
                 raw_paragraphs: list[tuple[str, str]] | None = None) -> bytes:
    """Build a .docx in memory. Either pass headings_and_bodies (level, heading,
    body_lines_or_None) for the common case, or raw_paragraphs (style_name,
    text) for edge cases the level-based helper can't express."""
    doc = Document()
    if raw_paragraphs is not None:
        for style_name, text in raw_paragraphs:
            p = doc.add_paragraph(text)
            p.style = doc.styles[style_name]
    else:
        for level, heading, body in headings_and_bodies or []:
            doc.add_heading(heading, level=level)
            for line in body or []:
                doc.add_paragraph(line)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_heading_1_and_2_both_work_identically():
    data = _build_docx(headings_and_bodies=[
        (1, "How can I reschedule my exam?", ["Send us a request and we'll review it."]),
        (2, "I missed my test window", ["Whether it can be retaken depends on the client."]),
    ])
    markdown, warnings = docx_to_markdown(data)

    assert warnings == []
    assert markdown == (
        "# How can I reschedule my exam?\n"
        "Send us a request and we'll review it.\n"
        "## I missed my test window\n"
        "Whether it can be retaken depends on the client."
    )


def test_statement_style_heading_is_not_required_to_be_a_question():
    data = _build_docx(headings_and_bodies=[
        (1, "I forgot my password", ["Use the reset link on the login page."]),
    ])
    markdown, warnings = docx_to_markdown(data)
    assert warnings == []
    assert "I forgot my password" in markdown


def test_heading_with_no_answer_is_dropped_and_warned():
    """A Heading 1 used as a bare section title, with no answer before the
    next heading, must not become a near-empty junk entry in the KB."""
    data = _build_docx(headings_and_bodies=[
        (1, "Scheduling", None),  # section title, no body
        (2, "How can I reschedule my exam?", ["Send us a request."]),
    ])
    markdown, warnings = docx_to_markdown(data)

    assert "Scheduling" not in markdown
    assert "How can I reschedule my exam?" in markdown
    assert len(warnings) == 1
    assert "Scheduling" in warnings[0]
    assert "dropped" in warnings[0]


def test_trailing_heading_with_no_answer_is_dropped():
    data = _build_docx(headings_and_bodies=[
        (1, "How can I reschedule my exam?", ["Send us a request."]),
        (1, "Contact us", None),  # last heading, nothing follows it either
    ])
    markdown, warnings = docx_to_markdown(data)
    assert "Contact us" not in markdown
    assert len(warnings) == 1


def test_heading_3_is_treated_as_body_text_not_a_new_question():
    """Only Heading 1/2 are recognized as question boundaries — this is a
    deliberate simplicity choice, not an oversight."""
    data = _build_docx(headings_and_bodies=[
        (1, "How can I reschedule my exam?", ["Send us a request."]),
    ])
    doc = Document(io.BytesIO(data))
    doc.add_heading("A stray sub-point", level=3)
    doc.add_paragraph("More detail here.")
    buf = io.BytesIO()
    doc.save(buf)

    markdown, warnings = docx_to_markdown(buf.getvalue())

    assert warnings == []
    # The level-3 text is folded into the SAME question's answer, not split out.
    assert "A stray sub-point" in markdown
    assert markdown.count("#") == 1  # only the one real Heading 1 line


def test_body_text_before_any_heading_is_dropped():
    data = _build_docx(raw_paragraphs=[
        ("Normal", "Welcome to the FAQ."),
        ("Heading 1", "How can I reschedule my exam?"),
        ("Normal", "Send us a request."),
    ])
    markdown, warnings = docx_to_markdown(data)
    assert "Welcome to the FAQ." not in markdown
    assert "How can I reschedule my exam?" in markdown


def test_empty_heading_text_is_skipped():
    data = _build_docx(raw_paragraphs=[
        ("Heading 1", ""),  # someone applied the style and typed nothing
        ("Heading 1", "How can I reschedule my exam?"),
        ("Normal", "Send us a request."),
    ])
    markdown, warnings = docx_to_markdown(data)
    assert "How can I reschedule my exam?" in markdown


def test_document_with_no_headings_warns():
    data = _build_docx(raw_paragraphs=[("Normal", "Just some plain text.")])
    markdown, warnings = docx_to_markdown(data)
    assert markdown == ""
    assert any("No Heading 1 / Heading 2" in w for w in warnings)


def test_not_a_docx_file_raises_a_clear_error():
    with pytest.raises(ValueError, match=r"\.docx"):
        docx_to_markdown(b"this is not a real docx file")


def test_output_feeds_chunk_markdown_correctly():
    """The whole point: chunk_markdown must not need to change at all."""
    data = _build_docx(headings_and_bodies=[
        (1, "How can I reschedule my exam?", ["Send us a request."]),
        (2, "I missed my test window", ["Depends on the client."]),
    ])
    markdown, warnings = docx_to_markdown(data)
    assert warnings == []

    chunks = chunk_markdown(markdown)
    assert len(chunks) == 2
    assert chunks[0] == "# How can I reschedule my exam?\nSend us a request."
    assert chunks[1] == "## I missed my test window\nDepends on the client."
